from pathlib import Path
import openpyxl,copy,json
p=Path(__file__).resolve().parent
a=openpyxl.load_workbook(p/'before_ratio.xlsx',data_only=False)
b=openpyxl.load_workbook(p/'Galicia_options_analysis.xlsx',data_only=False)
changes=[];styles=[];features=[]
for s in a:
    t=b[s.title]
    for row in s:
        for c in row:
            other=t[c.coordinate]
            if c.value!=other.value:changes.append((s.title,c.coordinate))
            if any(copy.copy(getattr(c,k))!=copy.copy(getattr(other,k)) for k in ['font','fill','alignment','border','protection']) or c.number_format!=other.number_format:styles.append((s.title,c.coordinate))
    if len(s._charts)!=len(t._charts) or len(s.tables)!=len(t.tables):features.append(s.title)
print('Prior-sheet value/formula differences:',len(changes),changes[:5])
print('Prior-sheet style differences:',len(styles),styles[:5])
print('Prior-sheet chart/table count differences:',features)
c=openpyxl.load_workbook(p/'Galicia_options_analysis.xlsx',data_only=True)
errors=[(s.title,x.coordinate,x.value) for s in c for row in s for x in row if x.data_type=='e' and not(s.title=='Ratio explorer' and x.column in [20,21,22] and x.value=='#N/A')]
print('Unexpected errors:',errors[:5])
print(json.dumps([(cell,c['Ratio explorer'][cell].value,str(b['Ratio explorer'][cell].alignment)) for cell in ['D12','A14','A15']]))
assert not changes and not styles and not features and not errors
print('Workbook preservation verified.')
