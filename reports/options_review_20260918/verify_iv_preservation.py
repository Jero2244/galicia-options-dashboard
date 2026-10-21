from pathlib import Path
import openpyxl,copy,zipfile,xml.etree.ElementTree as ET
p=Path(__file__).resolve().parent
a=openpyxl.load_workbook(p/'before_iv_curve.xlsx');b=openpyxl.load_workbook(p/'Galicia_options_analysis.xlsx')
cache={}; changes=[];styles=[]
for s in a:
    t=b[s.title]
    for row in s:
        for c in row:
            other=t[c.coordinate]
            if c.value!=other.value:changes.append((s.title,c.coordinate))
            key=(c.style_id,other.style_id)
            if key not in cache:cache[key]=all(copy.copy(getattr(c,k))==copy.copy(getattr(other,k)) for k in ['font','fill','alignment','border','protection']) and c.number_format==other.number_format
            if not cache[key]:styles.append((s.title,c.coordinate))
    assert len(s._charts)==len(t._charts) and len(s.tables)==len(t.tables)
    assert str(s.merged_cells)==str(t.merged_cells)
print('Original sheet value/formula changes:',len(changes),'Style changes:',len(styles))
assert not changes and not styles
with zipfile.ZipFile(p/'Galicia_options_analysis.xlsx') as z:
    for name in z.namelist():
        if '/charts/chart' in name and name.endswith('.xml'):
            xml=z.read(name).decode()
            if 'scatterChart' in xml:
                root=ET.fromstring(xml);ns={'c':'http://schemas.openxmlformats.org/drawingml/2006/chart'}
                print('Scatter style:',root.find('.//c:scatterStyle',ns).attrib)
                print('XY bindings:',[x.text for x in root.findall('.//c:xVal//c:f',ns)])
print('Preservation passed.')
